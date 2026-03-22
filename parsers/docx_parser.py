import zipfile
import re
from .base_parser import BaseParser

try:
    import docx
except ImportError:
    raise ImportError("缺少依赖库。请在终端运行: pip install python-docx")

class DocxParser(BaseParser):
    """
    处理 .docx 格式文档的解析器，自带“伪造后缀”的降级解析容错机制。
    """

    def extract_text(self, file_path: str) -> str:
        super().extract_text(file_path)
        
        try:
            doc = docx.Document(file_path)
            full_text = []
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if text: full_text.append(text)
                    
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data: full_text.append(" | ".join(row_data))
                        
            return "\n".join(full_text)
            
        except zipfile.BadZipFile:
            # 【容错降级核心】：文件不是合法的压缩包(docx本质是zip)
            # 极大概率是学生把 .txt 或 .html 强行改了后缀名。我们尝试强行提取纯文本。
            try:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    raw_content = f.read()
                    # 粗暴地剔除可能存在的 HTML 标签
                    clean_text = re.sub(r'<[^>]+>', ' ', raw_content).strip()
                    if clean_text:
                        return clean_text
            except Exception:
                pass
            raise ValueError("文件底层损坏或为旧版二进制 .doc，请让用户另存为标准 .docx 格式")
            
        except Exception as e:
            raise Exception(f"读取 DOCX 文件 {file_path} 时发生错误: {str(e)}")