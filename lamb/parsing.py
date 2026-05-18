"""Document parsing adapters used by LAMB workflows."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Iterable, List


SUPPORTED_EXTENSIONS = {
    ".txt": "Plain text",
    ".md": "Markdown",
    ".markdown": "Markdown",
    ".csv": "CSV",
    ".json": "JSON",
    ".docx": "Word document",
    ".pdf": "PDF",
}


class ParseError(RuntimeError):
    """Raised when a document cannot be parsed into text."""


def is_supported_extension(extension: str) -> bool:
    return extension.lower() in SUPPORTED_EXTENSIONS


def supported_extensions() -> List[str]:
    return sorted(SUPPORTED_EXTENSIONS)


def parse_document(path: str | Path) -> str:
    """Extract plain text from a supported document."""

    file_path = Path(path)
    extension = file_path.suffix.lower()
    if extension in {".txt", ".md", ".markdown"}:
        return _read_text(file_path)
    if extension == ".csv":
        return _read_csv(file_path)
    if extension == ".json":
        return _read_json(file_path)
    if extension == ".docx":
        return _read_docx(file_path)
    if extension == ".pdf":
        return _read_pdf(file_path)
    raise ParseError(f"unsupported file extension: {extension}")


def _read_text(path: Path) -> str:
    encodings = ("utf-8-sig", "utf-8", "gbk")
    errors: List[str] = []
    for encoding in encodings:
        try:
            return path.read_text(encoding=encoding).strip()
        except UnicodeDecodeError as exc:
            errors.append(f"{encoding}: {exc}")
    raise ParseError(f"failed to decode text file {path}: {'; '.join(errors)}")


def _read_csv(path: Path, row_limit: int = 5000) -> str:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(4096)
            handle.seek(0)
            dialect = csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
            reader = csv.reader(handle, dialect)
            rows = []
            for index, row in enumerate(reader):
                if index >= row_limit:
                    rows.append([f"... truncated after {row_limit} rows ..."])
                    break
                rows.append([cell.strip() for cell in row])
    except UnicodeDecodeError:
        with path.open("r", encoding="gbk", newline="") as handle:
            reader = csv.reader(handle)
            rows = [[cell.strip() for cell in row] for row in reader]
    except Exception as exc:
        raise ParseError(f"failed to read CSV {path}: {exc}") from exc

    return _tabulate_rows(rows)


def _read_json(path: Path) -> str:
    try:
        data = json.loads(_read_text(path))
    except json.JSONDecodeError as exc:
        raise ParseError(f"failed to decode JSON {path}: {exc}") from exc
    return json.dumps(data, ensure_ascii=False, indent=2)


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ParseError("python-docx is required to parse .docx files") from exc

    try:
        document = docx.Document(str(path))
        parts: List[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()
    except Exception as exc:
        raise ParseError(f"failed to parse DOCX {path}: {exc}") from exc


def _read_pdf(path: Path) -> str:
    try:
        import fitz
    except ImportError as exc:
        raise ParseError("pymupdf is required to parse .pdf files") from exc

    try:
        pages = []
        with fitz.open(str(path)) as document:
            for page_index in range(len(document)):
                text = document.load_page(page_index).get_text().strip()
                if text:
                    pages.append(text)
        return "\n\n".join(pages).strip()
    except Exception as exc:
        raise ParseError(f"failed to parse PDF {path}: {exc}") from exc


def _tabulate_rows(rows: Iterable[Iterable[str]]) -> str:
    lines = []
    for row in rows:
        lines.append(" | ".join(str(cell) for cell in row))
    return "\n".join(lines).strip()
